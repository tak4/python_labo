from difflib import SequenceMatcher
import openpyxl


def check_similarity_threshold(a: str, b: str, ratio_threshold: float = 0.5) -> tuple[bool, float]:
    over_threshold = False
    ratio_result = SequenceMatcher(None, a, b).ratio()
    if ratio_result >= ratio_threshold:
        over_threshold = True
    return (over_threshold, ratio_result)


wb = openpyxl.load_workbook("check_target.xlsx") # Excelファイルの読み込み

sheet_names = wb.sheetnames # シート名一覧取得

ws_target = wb['target'] # シートの取得
ws_list = wb['list'] # シートの取得

results = {}

row_target = 1
while True:
    value_target = ws_target.cell(row=row_target, column=1).value
    if value_target == None:
        break

    max_threshold = 0

    row_list = 1
    while True:
        value_list = ws_list.cell(row=row_list, column=1).value
        if value_list == None:
            break

        value_target = value_target.rstrip()
        value_list = value_list.rstrip()

        over_threshold, threshold = check_similarity_threshold(value_target, value_list)
        if max_threshold < threshold:
            max_threshold = threshold
            results[value_target] = (value_list, threshold)

        row_list += 1

    row_target += 1

for r in results:
    print(r)