import argparse
from difflib import SequenceMatcher
import openpyxl
from pathlib import Path


def check_similarity_threshold(a: str, b: str, ratio_threshold: float = 0.5) -> tuple[bool, float]:
    over_threshold = False
    ratio_result = SequenceMatcher(None, a, b).ratio()
    if ratio_result >= ratio_threshold:
        over_threshold = True
    return (over_threshold, ratio_result)

def check_by_difflib(
        criteria_wb_name: str, 
        criteria_ws_name: str, 
        criteria_index_col: int, 
        criteria_data_col: int,
        target_wb_name: str, 
        target_ws_name: str, 
        target_index_col: int, 
        target_data_col: int,
        output_dir: str
        ):

    path_criteria_wb = Path(criteria_wb_name).resolve()
    target_wb = Path(target_wb_name).resolve()
    output_dir_path = Path(output_dir).resolve()
    output_dir_path.mkdir(exist_ok=True)
    output_path = output_dir_path / f"result_{path_criteria_wb.name}"

    criteria_wb = openpyxl.load_workbook(path_criteria_wb)
    target_wb = openpyxl.load_workbook(target_wb)

    criteria_ws = criteria_wb[criteria_ws_name]
    target_ws = target_wb[target_ws_name]

    results = []

    criteria_ws_max_row = criteria_ws.max_row
    target_ws_max_row = target_ws.max_row
    for criteria_row in criteria_ws.iter_rows(min_row=1, max_row=criteria_ws_max_row, values_only=True):
        if criteria_row[criteria_data_col-1] is None:
            continue
        criteria_index_no = criteria_row[criteria_index_col-1]
        criteria_line = str(criteria_row[criteria_data_col-1]).strip()
        if not criteria_line:
            continue

        best_result = (criteria_index_no, criteria_line, None, None, 0.0, False)

        max_threshold = 0
        for target_row in target_ws.iter_rows(min_row=1, max_row=target_ws_max_row, values_only=True):
            if target_row[target_data_col-1] is None:
                continue
            target_index_no = target_row[target_index_col-1]
            target_line = str(target_row[target_data_col-1]).strip()
            if not target_line:
                continue

            over_threshold, threshold = check_similarity_threshold(criteria_line, target_line)
            if max_threshold < threshold:
                max_threshold = threshold
                best_result = (criteria_index_no, criteria_line, target_index_no, target_line, threshold, over_threshold)

        results.append(best_result)

    output_wb = openpyxl.Workbook()
    output_ws = output_wb.active
    output_ws.title = 'result'

    output_ws.append(['criteria_index_no', 'criteria_line', 'target_index_no', 'target_line', 'threshold', 'over_threshold'])
    for r in results:
        print(r)
        output_ws.append(r)

    # Excel ワークブック保存
    output_wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="")
    parser.add_argument("--criteria_wb", type=str, default="./pre_process/output/split_org_wb.xlsx", help="基準リストのワークブック名")
    parser.add_argument("--criteria_ws", type=str, default="split", help="基準リストのシート名")
    parser.add_argument("--criteria_index_col", type=int, default=1, help="基準リストの項番の列番号")
    parser.add_argument("--criteria_data_col", type=int, default=2, help="基準リストのデータの列番号")

    parser.add_argument("--target_wb", type=str, default="./pre_process/output/split_org_target_wb.xlsx", help="比較対象リストのワークブック名")
    parser.add_argument("--target_ws", type=str, default="split", help="比較対象リストのシート名")
    parser.add_argument("--target_index_col", type=int, default=1, help="比較対象リストの項番の列番号")
    parser.add_argument("--target_data_col", type=int, default=2, help="比較対象リストのデータの列番号")

    parser.add_argument("--output_dir", type=str, default="./process/output", help="出力先ディレクトリ")

    args = parser.parse_args()

    check_by_difflib(
        args.criteria_wb, 
        args.criteria_ws, 
        args.criteria_index_col, 
        args.criteria_data_col,
        args.target_wb, 
        args.target_ws, 
        args.target_index_col, 
        args.target_data_col,
        args.output_dir
    )


if __name__ == "__main__":
    main()
