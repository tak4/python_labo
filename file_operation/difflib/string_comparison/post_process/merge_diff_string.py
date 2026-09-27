import argparse
from difflib import SequenceMatcher
from html import escape
import openpyxl
from string import Template
from pathlib import Path

page_template = Template("""
<!DOCTYPE html>
<html lang="ja">
<head>
<meta charset="UTF-8">
<style>
body { font-family: sans-serif; }
.equal { background-color: #f2aaf8; }
.changed { background-color: #fff3a3; }
.deleted { background-color: #ffb3b3; text-decoration: line-through; }
.added { background-color: #b3ffb3; }
.comparison th {
    border-collapse: collapse;
    border-bottom: 1px solid #ddd;
}
.comparison td {
    border-collapse: collapse;
    border-bottom: 1px solid #ddd;
}
.comparison th:nth-child(1),
.comparison td:nth-child(1) {
    width: 20px;
}
.comparison th:nth-child(2),
.comparison td:nth-child(2) {
    width: 40px;
}
.comparison th:nth-child(3),
.comparison td:nth-child(3) {
    width: 600px;
}
.comparison th:nth-child(4),
.comparison td:nth-child(4) {
    width: 40px;
}
.comparison th:nth-child(5),
.comparison td:nth-child(5) {
    width: 600px;
}
.comparison th:nth-child(6),
.comparison td:nth-child(6) {
    width: 40px;
}
.comparison th:nth-child(7),
.comparison td:nth-child(7) {
    width: 40px;
}
.comparison th:nth-child(8),
.comparison td:nth-child(8) {
    width: 600px;
}
.comparison th:nth-child(9),
.comparison td:nth-child(9) {
    width: 40px;
}
.comparison th:nth-child(10),
.comparison td:nth-child(10) {
    width: 40px;
}
.comparison th:nth-child(11),
.comparison td:nth-child(11) {
    width: 600px;
}
.comparison th:nth-child(12),
.comparison td:nth-child(12) {
    width: 40px;
}
</style>
</head>
<body>
    <table>
        <tr class="comparison">
            <td>#</td>
            <td>idx_a</td>
            <td>wb_a</td>
            <td>idx_b</td>
            <td>wb_b</td>
            <td>ratio_b</td>
            <td>idx_c</td>
            <td>wb_c</td>
            <td>ratio_c</td>
            <td>idx_d</td>
            <td>wb_d</td>
            <td>ratio_d</td>
        </tr>
        $comparisons
    </table>
</body>
</html>
""")

def to_text(value):
    if value is None:
        return ""
    return str(value)

def create_pair_html(result_list: list, idx: int) -> str:
    index_a, text_a, index_b, text_b, ratio_b, index_c, text_c, ratio_c, index_d, text_d, ratio_d = result_list

    index_a = to_text(index_a)
    index_b = to_text(index_b)
    index_c = to_text(index_c)
    index_d = to_text(index_d)

    text_a = to_text(text_a)
    text_b = to_text(text_b)
    text_c = to_text(text_c)
    text_d = to_text(text_d)

    ratio_b = float(ratio_b) if ratio_b is not None else 0.0
    ratio_c = float(ratio_c) if ratio_c is not None else 0.0
    ratio_d = float(ratio_d) if ratio_d is not None else 0.0

    matcher_b = SequenceMatcher(None, text_a, text_b)
    matcher_c = SequenceMatcher(None, text_a, text_c)
    matcher_d = SequenceMatcher(None, text_a, text_d)

    text_a_b_html = []
    text_b_html = []
    text_a_c_html = []
    text_c_html = []
    text_a_d_html = []
    text_d_html = []

    # A vs B
    for tag, i1, i2, j1, j2 in matcher_b.get_opcodes():
        if tag == "equal":
            text = escape(text_a[i1:i2] or "")
            text_a_b_html.append(text)
            text_b_html.append(f'<span class="equal">{text}</span>')

        elif tag == "replace":
            text_a_part = escape(text_a[i1:i2] or "")
            text_b_part = escape(text_b[j1:j2] or "")
            # text_a_b_html.append(f'<span class="changed">{text_a_part}</span>')
            # text_b_html.append(f'<span class="changed">{text_b_part}</span>')
            text_a_b_html.append(text_a_part)
            text_b_html.append(text_b_part)

        elif tag == "delete":
            text = escape(text_a[i1:i2] or "")
            # text_a_b_html.append(f'<span class="deleted">{text}</span>')
            text_a_b_html.append(text)

        elif tag == "insert":
            text = escape(text_b[j1:j2] or "")
            # text_b_html.append(f'<span class="added">{text}</span>')
            text_b_html.append(text)

    # A vs C
    for tag, i1, i2, j1, j2 in matcher_c.get_opcodes():
        if tag == "equal":
            text = escape(text_a[i1:i2] or "")
            text_a_c_html.append(text)
            text_c_html.append(f'<span class="equal">{text}</span>')

        elif tag == "replace":
            text_a_part = escape(text_a[i1:i2] or "")
            text_c_part = escape(text_c[j1:j2] or "")
            # text_a_c_html.append(f'<span class="changed">{text_a_part}</span>')
            # text_c_html.append(f'<span class="changed">{text_c_part}</span>')
            text_a_c_html.append(text_a_part)
            text_c_html.append(text_c_part)

        elif tag == "delete":
            text = escape(text_a[i1:i2] or "")
            # text_a_c_html.append(f'<span class="deleted">{text}</span>')
            text_a_c_html.append(text)

        elif tag == "insert":
            text = escape(text_c[j1:j2] or "")
            # text_c_html.append(f'<span class="added">{text}</span>')
            text_c_html.append(text)

    # A vs D
    for tag, i1, i2, j1, j2 in matcher_d.get_opcodes():
        if tag == "equal":
            text = escape(text_a[i1:i2] or "")
            text_a_d_html.append(text)
            text_d_html.append(f'<span class="equal">{text}</span>')

        elif tag == "replace":
            text_a_part = escape(text_a[i1:i2] or "")
            text_d_part = escape(text_d[j1:j2] or "")
            # text_a_d_html.append(f'<span class="changed">{text_a_part}</span>')
            # text_d_html.append(f'<span class="changed">{text_d_part}</span>')
            text_a_d_html.append(text_a_part)
            text_d_html.append(text_d_part)

        elif tag == "delete":
            text = escape(text_a[i1:i2] or "")
            # text_a_d_html.append(f'<span class="deleted">{text}</span>')
            text_a_d_html.append(text)

        elif tag == "insert":
            text = escape(text_d[j1:j2] or "")
            # text_d_html.append(f'<span class="added">{text}</span>')
            text_d_html.append(text)

    return f"""
        <tr class="comparison">
            <td>{idx}</td>
            <td>{index_a}</td>
            <td>{''.join(text_a_b_html)}</td>
            <td>{index_b}</td>
            <td>{''.join(text_b_html)}</td>
            <td>{round(float(ratio_b), 2)}</td>
            <td>{index_c}</td>
            <td>{''.join(text_c_html)}</td>
            <td>{round(float(ratio_c), 2)}</td>
            <td>{index_d}</td>
            <td>{''.join(text_d_html)}</td>
            <td>{round(float(ratio_d), 2)}</td>
        </tr>
    """

def create_diff_page(comparison_list: list):
    sections = []
    for idx, result_list in enumerate(comparison_list, 1):
        sections.append(create_pair_html(result_list, idx))
    return page_template.substitute(comparisons="".join(sections))


def main():
    parser = argparse.ArgumentParser(description="")
    parser.add_argument("--merge_a_wb", type=str, default="./process/output_test/result_split_list_target_wb_a.xlsx", help="入力ワークブックA名")
    parser.add_argument("--merge_a_ws", type=str, default="result", help="入力ワークブックAのシート名")
    parser.add_argument("--merge_b_wb", type=str, default="./process/output_test/result_split_list_target_wb_b.xlsx", help="入力ワークブックB名")
    parser.add_argument("--merge_b_ws", type=str, default="result", help="入力ワークブックBのシート名")
    parser.add_argument("--merge_c_wb", type=str, default="./process/output_test/result_split_list_target_wb_c.xlsx", help="入力ワークブックC名")
    parser.add_argument("--merge_c_ws", type=str, default="result", help="入力ワークブックCのシート名")
    parser.add_argument("--merge_d_wb", type=str, default="./process/output_test/result_split_list_target_wb_d.xlsx", help="入力ワークブックD名")
    parser.add_argument("--merge_d_ws", type=str, default="result", help="入力ワークブックDのシート名")
    parser.add_argument("--output_dir", type=str, default="./post_process/output", help="出力先ディレクトリ")

    args = parser.parse_args()

    merge_a_wb_path = Path(args.merge_a_wb).resolve()
    merge_b_wb_path = Path(args.merge_b_wb).resolve()
    merge_c_wb_path = Path(args.merge_c_wb).resolve()
    merge_d_wb_path = Path(args.merge_d_wb).resolve()

    merge_a_wb = openpyxl.load_workbook(merge_a_wb_path, data_only=True)
    merge_b_wb = openpyxl.load_workbook(merge_b_wb_path, data_only=True)
    merge_c_wb = openpyxl.load_workbook(merge_c_wb_path, data_only=True)
    merge_d_wb = openpyxl.load_workbook(merge_d_wb_path, data_only=True)

    merge_a_ws = merge_a_wb[args.merge_a_ws]
    merge_b_ws = merge_b_wb[args.merge_b_ws]
    merge_c_ws = merge_c_wb[args.merge_c_ws]
    merge_d_ws = merge_d_wb[args.merge_d_ws]

    merge_a_ws_max_row = merge_a_ws.max_row

    merged_list = []
    for idx in range(2, merge_a_ws_max_row):
        merged_list.append(
            [cell.value for cell in merge_a_ws[idx][:]] +
            [cell.value for cell in merge_b_ws[idx][2:6]] +
            [cell.value for cell in merge_c_ws[idx][2:6]] +
            [cell.value for cell in merge_d_ws[idx][2:6]]
    )

    output_dir_path = Path(args.output_dir).resolve()
    output_dir_path.mkdir(exist_ok=True)
    output_path = output_dir_path / f"merged_{merge_a_wb_path.name}"

    output_wb = openpyxl.Workbook()
    output_ws = output_wb.active
    output_ws.title = 'merged'

    for r in merged_list:
        output_ws.append(r)

    # Excel ワークブック保存
    output_wb.save(output_path)

    comparison_list = []
    for row in merged_list:
        comparison_list.append([
            row[0],   # index_a
            row[1],   # text_a
            row[6],   # index_b
            row[7],   # text_b
            row[8],   # ratio_b
            row[10],  # index_c
            row[11],  # text_c
            row[12],  # ratio_c
            row[14],  # index_d
            row[15],  # text_d
            row[16],  # ratio_d
        ])

    page_html = create_diff_page(comparison_list)

    merged_html = output_dir_path / f"merged_{merge_a_wb_path.stem}.html"
    with open(merged_html, "w", encoding="utf-8") as f:
        f.write(page_html)

if __name__ == "__main__":
    main()
