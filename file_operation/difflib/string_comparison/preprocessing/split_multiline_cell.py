import argparse
import openpyxl
from pathlib import Path


def split_cell(input_wb_name: str, input_ws_name: str, input_index_col: int, input_data_col: int):

    input_wb_path = Path(input_wb_name).resolve()

    org_wb = openpyxl.load_workbook(input_wb_path, data_only=True) # Excelファイルの読み込み
    org_ws = org_wb[input_ws_name] # シートの取得

    split_cell_wb = openpyxl.Workbook()
    split_cell_ws = split_cell_wb.active
    split_cell_ws.title = 'split'

    max_row = org_ws.max_row
    for row in org_ws.iter_rows(min_row=2, max_row=max_row, values_only=True):
        if row[input_data_col-1] is None:
            continue
        row_no = row[input_index_col-1]
        lines = row[input_data_col-1].splitlines()
        for l in lines:
            line = l.strip()
            if line.startswith('='):
                line = "'" + line
            if len(line) > 0:
                split_cell_ws.append([row_no, line])

    # Excel ワークブック保存
    output_path = input_wb_path.with_name(f"split_{input_wb_path.name}")
    split_cell_wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="")
    parser.add_argument("input_wb", type=str, nargs="?", default="org_wb.xlsx", help="入力ワークブック名")
    parser.add_argument("input_ws", type=str, nargs="?", default="target", help="入力ワークブックのシート名")
    parser.add_argument("input_index_col", type=int, nargs="?", default=1, help="項番の列番号")
    parser.add_argument("input_data_col", type=int, nargs="?", default=2, help="データの列番号")

    args = parser.parse_args()

    split_cell(args.input_wb, args.input_ws, args.input_index_col, args.input_data_col)


if __name__ == "__main__":
    main()
